from datetime import date, timedelta
from decimal import Decimal
import json
import io

from django.contrib.auth.models import Group, User
from django.contrib.messages.storage.fallback import FallbackStorage
from django.contrib.sessions.middleware import SessionMiddleware
from django.http import HttpRequest
from django.test import RequestFactory, TestCase
from django.urls import reverse

from inventory.forms import ProductForm, ProductItemForm
from inventory.roles import ROLE_INVENTORY_MANAGER
from inventory.views import delete_products
from services.data_collection_2.create_withdrawal import create_withdrawal
from services.data_collection.barcode_resolution import resolve_product_from_barcode
from services.data_collection.data_collection import get_product_by_id, parse_barcode_data
from services.reporting.reporting import download_report
from services.data_storage.models import (
    Location,
    Product,
    ProductBarcodeAlias,
    ProductDeletionArchive,
    ProductIdentifier,
    ProductItem,
    StockRegistration,
    Supplier,
    Withdrawal,
)
from solutions.analytics.views import track_qc
from solutions.quality_control.models import QualityCheck
from solutions.quality_control.views import create_check
from openpyxl import load_workbook


class BarcodeParserTests(TestCase):
    def test_parse_3pr_includes_alias_fields(self):
        parsed = parse_barcode_data("3PR02814**0230105021**31.12.2026")
        self.assertEqual(parsed["barcode_type"], "3PR")
        self.assertEqual(parsed["alias_identifier_type"], "PARSED_PRODUCT_CODE")
        self.assertEqual(parsed["alias_identifier_value"], "3PR02814")
        self.assertEqual(parsed["lot_number"], "0230105021")
        self.assertEqual(parsed["expiry_date"], "31.12.2026")

    def test_parse_gs1_includes_gtin_alias(self):
        parsed = parse_barcode_data("01008476270080391727033110092625A")
        self.assertEqual(parsed["barcode_type"], "GS1")
        self.assertEqual(parsed["gtin"], "00847627008039")
        self.assertEqual(parsed["alias_identifier_type"], "GTIN")
        self.assertEqual(parsed["alias_identifier_value"], "00847627008039")
        self.assertEqual(parsed["lot_number"], "092625A")
        self.assertEqual(parsed["expiry_date"], "31.03.2027")

    def test_parse_gs1_with_lot_then_expiry_then_ai_240(self):
        parsed = parse_barcode_data("010401563098474910N15097172702072400671866300192860-09911250827")
        self.assertEqual(parsed["barcode_type"], "GS1")
        self.assertEqual(parsed["gtin"], "04015630984749")
        self.assertEqual(parsed["lot_number"], "N15097")
        self.assertEqual(parsed["expiry_date"], "07.02.2027")
        self.assertEqual(parsed["additional_product_id"], "0671866300192860-09911250827")

    def test_parse_live_ultraview_style_barcode(self):
        parsed = parse_barcode_data("010401563097217310N15098172706102400526980600192760-50011250917")
        self.assertEqual(parsed["gtin"], "04015630972173")
        self.assertEqual(parsed["lot_number"], "N15098")
        self.assertEqual(parsed["expiry_date"], "10.06.2027")
        self.assertEqual(parsed["additional_product_id"], "0526980600192760-50011250917")


class BarcodeResolverTests(TestCase):
    def setUp(self):
        self.supplier = Supplier.objects.create(name="Biocare")

    def test_resolves_from_barcode_alias_gtin(self):
        product = Product.objects.create(
            name="NKX3.1",
            supplier="THIRD_PARTY",
            supplier_ref=self.supplier,
            threshold=1,
            lead_time=timedelta(days=1),
        )
        ProductBarcodeAlias.objects.create(
            product=product,
            barcode_type="GS1",
            identifier_type="GTIN",
            identifier_value="00847627008039",
            is_active=True,
        )

        parsed = parse_barcode_data("01008476270080391727033110092625A")
        resolution = resolve_product_from_barcode(parsed, "01008476270080391727033110092625A")
        self.assertEqual(resolution["product"].id, product.id)
        self.assertEqual(resolution["source"], "barcode_alias_gtin")

    def test_falls_back_to_legacy_product_code(self):
        product = Product.objects.create(
            product_code="00660",
            name="Legacy Product",
            supplier="LEICA",
            threshold=1,
            lead_time=timedelta(days=1),
        )
        parsed = {
            "product_code": "00660",
            "raw_product_code": "00660",
            "normalized_product_code": "660",
            "gtin": "",
        }
        resolution = resolve_product_from_barcode(parsed, "00660")
        self.assertEqual(resolution["product"].id, product.id)
        self.assertEqual(resolution["source"], "legacy_product_code")


class BarcodeEndpointsTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_user(username="manager", password="pass123")
        group, _ = Group.objects.get_or_create(name=ROLE_INVENTORY_MANAGER)
        self.user.groups.add(group)

        self.supplier = Supplier.objects.create(name="Abcam")
        self.product = Product.objects.create(
            name="Amyloid A vial 1",
            supplier="THIRD_PARTY",
            supplier_ref=self.supplier,
            threshold=1,
            lead_time=timedelta(days=2),
        )

    def test_unresolved_barcode_returns_mapping_required(self):
        response = self.client.get(
            reverse("data:get_product_by_barcode"),
            {"barcode": "UNKNOWN-CODE-XYZ"},
        )
        self.assertEqual(response.status_code, 404)
        payload = response.json()
        self.assertTrue(payload["mapping_required"])
        self.assertFalse(payload["success"])

    def test_mapping_create_and_conflict_validation(self):
        self.client.force_login(self.user)
        response = self.client.post(
            reverse("data:create_barcode_mapping"),
            data={
                "product_id": str(self.product.uuid),
                "barcode_type": "GS1",
                "identifier_type": "GTIN",
                "identifier_value": "00847627008039",
                "raw_barcode": "01008476270080391727033110092625A",
            },
        )
        self.assertEqual(response.status_code, 200)
        self.assertTrue(response.json()["success"])
        self.assertTrue(
            ProductBarcodeAlias.objects.filter(
                identifier_type="GTIN",
                identifier_value="00847627008039",
                product=self.product,
            ).exists()
        )
        self.assertTrue(
            ProductIdentifier.objects.filter(
                identifier_type="GTIN",
                identifier_value="00847627008039",
                product=self.product,
            ).exists()
        )

        other = Product.objects.create(
            name="Other Product",
            supplier="THIRD_PARTY",
            threshold=1,
            lead_time=timedelta(days=1),
        )
        conflict = self.client.post(
            reverse("data:create_barcode_mapping"),
            data={
                "product_id": str(other.uuid),
                "barcode_type": "GS1",
                "identifier_type": "GTIN",
                "identifier_value": "00847627008039",
            },
        )
        self.assertEqual(conflict.status_code, 409)

    def test_dropdown_returns_uuid_values(self):
        self.client.force_login(self.user)
        response = self.client.get(reverse("data:product_dropdown"), {"q": "Amyloid"})
        self.assertEqual(response.status_code, 200)
        results = response.json()["results"]
        self.assertTrue(any(row["id"] == str(self.product.uuid) for row in results))


class WorkflowTests(TestCase):
    def setUp(self):
        self.factory = RequestFactory()
        self.user = User.objects.create_user(username="manager", password="pass123")
        group, _ = Group.objects.get_or_create(name=ROLE_INVENTORY_MANAGER)
        self.user.groups.add(group)

        self.product = Product.objects.create(
            name="NKX3.1",
            supplier="THIRD_PARTY",
            threshold=1,
            lead_time=timedelta(days=1),
        )
        ProductBarcodeAlias.objects.create(
            product=self.product,
            barcode_type="GS1",
            identifier_type="GTIN",
            identifier_value="00847627008039",
            is_active=True,
        )
        self.item = ProductItem.objects.create(
            product=self.product,
            lot_number="092625A",
            expiry_date=date(2027, 3, 31),
            current_stock=Decimal("5.00"),
            units_per_quantity=1,
        )
        self.partial_item = ProductItem.objects.create(
            product=self.product,
            lot_number="PART001",
            expiry_date=date(2027, 6, 1),
            current_stock=Decimal("5.00"),
            units_per_quantity=4,
            accumulated_partial=1,
            product_feature="unit",
        )
        self.volume_item = ProductItem.objects.create(
            product=self.product,
            lot_number="VOL001",
            expiry_date=date(2027, 7, 1),
            current_stock=Decimal("7.50"),
            units_per_quantity=Decimal("0.20"),
            product_feature="volume",
        )
        self.expired_item = ProductItem.objects.create(
            product=self.product,
            lot_number="EXP001",
            expiry_date=date(2025, 1, 1),
            current_stock=Decimal("10.00"),
            units_per_quantity=1,
        )
        self.barcode = "01008476270080391727033110092625A"
        self.expired_barcode = "01008476270080391725010110EXP001"

    def _build_post_request(self, path, data):
        request = self.factory.post(path, data=data)
        request.user = self.user
        middleware = SessionMiddleware(lambda req: None)
        middleware.process_request(request)
        request.session.save()
        setattr(request, "_messages", FallbackStorage(request))
        return request

    def _build_get_request(self, path, data=None):
        request = self.factory.get(path, data=data or {})
        request.user = self.user
        middleware = SessionMiddleware(lambda req: None)
        middleware.process_request(request)
        request.session.save()
        setattr(request, "_messages", FallbackStorage(request))
        return request

    def test_available_stock_excludes_expired_lots(self):
        self.assertEqual(self.product.get_available_stock(), Decimal("17.50"))
        self.assertEqual(self.product.get_full_items_in_stock(), 17)

    def test_register_stock_with_mapped_barcode(self):
        self.client.force_login(self.user)
        response = self.client.post(
            reverse("data_collection_3:register_stock"),
            data={"barcode": self.barcode},
        )
        self.assertEqual(response.status_code, 302)
        self.item.refresh_from_db()
        self.assertEqual(self.item.current_stock, Decimal("6.00"))
        registration = StockRegistration.objects.latest("id")
        self.assertEqual(str(registration.resolved_product_uuid), str(self.product.uuid))
        self.assertEqual(registration.barcode, self.barcode)

    def test_withdrawal_with_mapped_barcode(self):
        self.client.force_login(self.user)
        response = self.client.post(
            reverse("data_collection_2:create_withdrawal"),
            data={
                "barcode": self.barcode,
                "quantity": "1",
                "withdrawal_type": "unit",
                "withdrawal_mode": "full",
                "parts_withdrawn": "0",
                "lot_number": "092625A",
                "expiry_date": "31.03.2027",
            },
        )
        self.assertEqual(response.status_code, 302)
        self.item.refresh_from_db()
        self.assertEqual(self.item.current_stock, Decimal("4.00"))
        withdrawal = Withdrawal.objects.latest("id")
        self.assertEqual(str(withdrawal.resolved_product_uuid), str(self.product.uuid))

    def test_barcode_lookup_flags_expired_lot_and_keeps_available_stock(self):
        response = self.client.get(
            reverse("data:get_product_by_barcode"),
            {"barcode": self.expired_barcode},
        )
        self.assertEqual(response.status_code, 200)
        payload = response.json()
        self.assertTrue(payload["expired_lot"])
        self.assertEqual(payload["stock"], "10.00")
        self.assertEqual(payload["current_stock"], "17.50")

    def test_manual_product_lookup_marks_lot_requiring_qc_action(self):
        QualityCheck.objects.create(
            product_item=self.item,
            performed_by=self.user,
            status=QualityCheck.STATUS_COMPLETED,
            result="pass",
        )
        request = self._build_get_request("/data/get-product-by-id/", {"id": str(self.product.id)})
        response = get_product_by_id(request)
        self.assertEqual(response.status_code, 200)
        payload = json.loads(response.content)
        lots_by_number = {row["lot_number"]: row for row in payload["lots"]}
        self.assertFalse(lots_by_number["092625A"]["qc_action_required"])
        self.assertTrue(lots_by_number["PART001"]["qc_action_required"])
        self.assertTrue(lots_by_number["092625A"]["qc_passed"])

    def test_qc_create_check_preselects_requested_product_item(self):
        request = self._build_get_request(
            "/quality-control/checks/create/",
            {"product_item_id": str(self.partial_item.id)},
        )
        response = create_check(request)
        self.assertEqual(response.status_code, 200)
        self.assertIn(str(self.partial_item.id), response.content.decode("utf-8"))

    def test_qc_create_check_renders_product_and_lot_lookup_inputs(self):
        request = self._build_get_request("/quality-control/checks/create/")
        response = create_check(request)
        self.assertEqual(response.status_code, 200)
        content = response.content.decode("utf-8")
        self.assertIn('name="product_name_lookup"', content)
        self.assertIn('name="lot_number_lookup"', content)
        self.assertIn(f'data-product-name="{self.product.name}"', content)
        self.assertIn(f'data-lot-number="{self.partial_item.lot_number}"', content)

    def test_qc_create_check_accepts_lot_lookup_without_dropdown_selection(self):
        self.client.force_login(self.user)
        response = self.client.post(
            reverse("quality_control:create_check"),
            data={
                "product_item": "",
                "product_name_lookup": self.product.name,
                "lot_number_lookup": self.partial_item.lot_number,
                "status": QualityCheck.STATUS_COMPLETED,
                "result": "pass",
                "test_reference": "QC-ALT-1",
                "notes": "Resolved from lookup fields",
            },
        )
        self.assertEqual(response.status_code, 302)
        qc = QualityCheck.objects.latest("id")
        self.assertEqual(qc.product_item_id, self.partial_item.id)
        self.assertEqual(qc.test_reference, "QC-ALT-1")

    def test_qc_create_check_requires_disambiguation_for_product_name_only(self):
        self.client.force_login(self.user)
        response = self.client.post(
            reverse("quality_control:create_check"),
            data={
                "product_item": "",
                "product_name_lookup": self.product.name,
                "lot_number_lookup": "",
                "status": QualityCheck.STATUS_PENDING,
                "result": "",
                "test_reference": "QC-AMBIG",
                "notes": "",
            },
        )
        self.assertEqual(response.status_code, 200)
        self.assertContains(
            response,
            "Multiple lots match that selection. Add the lot number or use the product lot dropdown.",
        )

    def test_delete_product_archives_snapshot_and_preserves_history(self):
        ProductIdentifier.objects.create(
            product=self.product,
            identifier_type="NAME",
            identifier_value="NKX3.1 Main",
        )
        Withdrawal.objects.create(
            product_item=self.item,
            quantity=Decimal("1.00"),
            withdrawal_type="unit",
            user=self.user,
        )
        StockRegistration.objects.create(
            product_item=self.item,
            quantity=1,
            user=self.user,
        )

        request = self._build_post_request(
            "/inventory/delete_products/",
            data={
                "product_id": str(self.product.id),
                "password": "pass123",
            },
        )
        response = delete_products(request)
        self.assertEqual(response.status_code, 302)
        self.assertFalse(Product.objects.filter(id=self.product.id).exists())

        archive = ProductDeletionArchive.objects.get(deleted_product_name="NKX3.1")
        self.assertEqual(archive.deleted_by, self.user)
        self.assertEqual(archive.deleted_product_uuid, self.product.uuid)
        self.assertTrue(any(item["lot_number"] == "092625A" for item in archive.items_snapshot))
        self.assertTrue(any(identifier["identifier_value"] == "NKX3.1 Main" for identifier in archive.identifiers_snapshot))

        withdrawal = Withdrawal.objects.latest("id")
        registration = StockRegistration.objects.latest("id")
        self.assertEqual(withdrawal.product_name, "NKX3.1")
        self.assertEqual(registration.product_name, "NKX3.1")

    def test_track_qc_filters_by_result_lot_and_product(self):
        QualityCheck.objects.create(
            product_item=self.item,
            performed_by=self.user,
            status=QualityCheck.STATUS_COMPLETED,
            result="pass",
        )
        QualityCheck.objects.create(
            product_item=self.partial_item,
            performed_by=self.user,
            status=QualityCheck.STATUS_COMPLETED,
            result="fail",
        )

        request = self._build_get_request(
            "/analytics/track-qc/",
            {
                "result": "fail",
                "product_name": "NKX3.1",
                "lot_number": "PART",
            },
        )
        response = track_qc(request)
        self.assertEqual(response.status_code, 200)
        content = response.content.decode("utf-8")
        self.assertIn("PART001", content)
        self.assertNotIn("092625A", content)

    def test_track_qc_available_in_reports_page_for_inventory_manager(self):
        request = self._build_get_request("/analytics/track-qc/")
        response = track_qc(request)
        self.assertEqual(response.status_code, 200)
        self.assertIn("Track QC", response.content.decode("utf-8"))

    def test_download_report_preview_is_paginated_to_100_rows(self):
        for index in range(101):
            StockRegistration.objects.create(
                product_item=self.item,
                quantity=1,
                user=self.user,
                product_code=f"CODE-{index}",
                product_name=f"Product {index}",
                lot_number=f"LOT-{index}",
                expiry_date=date(2027, 1, 1),
            )

        request = self._build_get_request("/analytics/reports/download/")
        response = download_report(request)
        self.assertEqual(response.status_code, 200)
        content = response.content.decode("utf-8")
        self.assertIn("Page 1 of 2", content)
        self.assertIn("Showing 1-100 of 101 records.", content)
        self.assertIn("No. of Products in Stock", content)

    def test_download_report_csv_contains_requested_columns(self):
        QualityCheck.objects.create(
            product_item=self.item,
            performed_by=self.user,
            status=QualityCheck.STATUS_COMPLETED,
            result="pass",
        )
        Withdrawal.objects.create(
            product_item=self.item,
            quantity=Decimal("1.00"),
            withdrawal_type="unit",
            user=self.user,
        )
        StockRegistration.objects.create(
            product_item=self.item,
            quantity=2,
            user=self.user,
        )

        request = self._build_get_request(
            "/analytics/reports/download/",
            {"download": "csv"},
        )
        response = download_report(request)
        self.assertEqual(response.status_code, 200)
        csv_text = response.content.decode("utf-8")
        self.assertIn("Date (Registration),Product,Product Code,Lot Number,Expire Date,Quantity,No. of Products in Stock,Location,User (Registered),User (Withdraw),Date Withdraw,QC Status,Person QC'ed,Date QC'ed", csv_text)
        self.assertIn("Passed", csv_text)

    def test_download_report_excel_includes_product_summary_sheet(self):
        StockRegistration.objects.create(
            product_item=self.item,
            quantity=2,
            user=self.user,
        )

        request = self._build_get_request(
            "/analytics/reports/download/",
            {"download": "excel"},
        )
        response = download_report(request)
        self.assertEqual(response.status_code, 200)

        workbook = load_workbook(io.BytesIO(response.content))
        self.assertIn("Inventory Report", workbook.sheetnames)
        self.assertIn("Product Summary", workbook.sheetnames)

        summary_sheet = workbook["Product Summary"]
        headers = [summary_sheet.cell(row=1, column=idx).value for idx in range(1, 5)]
        self.assertEqual(headers, ["Product", "Product Code", "Items In Stock", "Expired Lots"])
        values = [summary_sheet.cell(row=2, column=idx).value for idx in range(1, 5)]
        self.assertEqual(values[0], "NKX3.1")
        self.assertEqual(values[2], "17.50")
        self.assertIn("EXP001", values[3])

    def test_stock_admin_lookup_uses_alias_resolution(self):
        self.client.force_login(self.user)
        response = self.client.get(
            reverse("data_collection_1:stock_admin"),
            {"raw": self.barcode},
        )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.context["editing_product"].id, self.product.id)

    def test_stock_admin_volume_entry_derives_internal_lot_fields(self):
        form = ProductItemForm(
            data={
                "lot_number": "LOT-REAG-1",
                "expiry_date": "2027-06-30",
                "product_feature": "volume",
                "full_volume": "7.50",
                "partial_withdrawal_volume": "0.20",
                "current_stock": "",
                "units_per_quantity": "",
                "accumulated_partial": "",
            }
        )
        self.assertTrue(form.is_valid(), form.errors)
        item = form.save(commit=False)
        self.assertEqual(item.product_feature, "volume")
        self.assertEqual(item.current_stock, Decimal("7.50"))
        self.assertEqual(item.units_per_quantity, Decimal("0.20"))
        self.assertEqual(item.accumulated_partial, 0)

    def test_manual_unit_partial_withdrawal_updates_partial_tracking(self):
        request = self._build_post_request(
            "/stock/create_withdrawal/",
            data={
                "product_dropdown": str(self.product.id),
                "manual_lot_item_id": str(self.partial_item.id),
                "barcode_manual": self.product.product_code or "",
                "withdrawal_type": "part",
                "withdrawal_mode": "part",
                "parts_withdrawn": "3",
            },
        )
        response = create_withdrawal(request)
        self.assertEqual(response.status_code, 302)
        self.partial_item.refresh_from_db()
        self.assertEqual(self.partial_item.current_stock, Decimal("4.00"))
        self.assertEqual(self.partial_item.accumulated_partial, 0)
        withdrawal = Withdrawal.objects.latest("id")
        self.assertEqual(withdrawal.withdrawal_type, "part")
        self.assertEqual(withdrawal.quantity, Decimal("1.00"))
        self.assertEqual(withdrawal.parts_withdrawn, 3)

    def test_manual_volume_partial_withdrawal_updates_remaining_volume(self):
        request = self._build_post_request(
            "/stock/create_withdrawal/",
            data={
                "product_dropdown": str(self.product.id),
                "manual_lot_item_id": str(self.volume_item.id),
                "barcode_manual": self.product.product_code or "",
                "withdrawal_type": "part",
                "withdrawal_mode": "part",
                "parts_withdrawn": "3",
            },
        )
        response = create_withdrawal(request)
        self.assertEqual(response.status_code, 302)
        self.volume_item.refresh_from_db()
        self.assertEqual(self.volume_item.current_stock, Decimal("6.90"))
        withdrawal = Withdrawal.objects.latest("id")
        self.assertEqual(withdrawal.withdrawal_type, "part")
        self.assertEqual(withdrawal.quantity, Decimal("0.60"))
        self.assertEqual(withdrawal.parts_withdrawn, 3)

    def test_stock_admin_location_choices_are_limited(self):
        Location.objects.create(name="Histology")
        form = ProductForm()
        location_names = list(form.fields["location"].queryset.values_list("name", flat=True))
        self.assertEqual(location_names, ["Ccentral", "Cold Store"])

    def test_qc_lot_status_lookup_uses_alias_resolution(self):
        self.client.force_login(self.user)
        response = self.client.get(
            reverse("quality_control:lot_status"),
            {"barcode": self.barcode},
        )
        self.assertEqual(response.status_code, 200)
        selected = response.context["selected_product"]
        self.assertIsNotNone(selected)
        self.assertEqual(selected.id, self.product.id)

    def test_delete_expired_lot_creates_discard_withdrawal_and_redirects(self):
        self.client.force_login(self.user)
        response = self.client.post(f"/stock/delete_lot/{self.expired_item.id}/")
        self.assertEqual(response.status_code, 302)
        self.assertFalse(ProductItem.objects.filter(id=self.expired_item.id).exists())
        withdrawal = Withdrawal.objects.latest("id")
        self.assertEqual(withdrawal.withdrawal_type, "discarded")
        self.assertEqual(withdrawal.lot_number, "EXP001")
