import csv
import datetime
import io
from collections import defaultdict
from decimal import Decimal

from django.core.paginator import Paginator
from django.http import HttpResponse
from django.shortcuts import render
from django.utils import timezone
from django.utils.timezone import make_aware, now
from openpyxl import Workbook

from services.data_storage.models import ProductItem, StockRegistration, Withdrawal

try:
    from solutions.quality_control.models import QualityCheck
except Exception:
    QualityCheck = None


# One row per reagent lot. Received/Withdrawn/Current stock reconcile
# (received - withdrawn ~= current), so a single withdrawal reads as 1, not "all".
LOT_SUMMARY_COLUMNS = [
    ("product_name", "Product"),
    ("product_code", "Product Code"),
    ("lot_number", "Lot Number"),
    ("expiry_date", "Expiry Date"),
    ("expiry_status", "Expiry Status"),
    ("location_name", "Location"),
    ("received", "Received (total)"),
    ("withdrawn", "Withdrawn (total)"),
    ("current_stock", "Current Stock"),
    ("qc_status", "QC Status"),
    ("qc_by", "QC By"),
    ("qc_date", "QC Date"),
    ("first_registered", "First Registered"),
    ("registered_by", "Registered By"),
    ("last_withdrawal", "Last Withdrawal"),
    ("withdrawn_by", "Last Withdrawn By"),
]

# One row per stock movement (each registration / withdrawal is its own event).
MOVEMENTS_COLUMNS = [
    ("movement_date", "Date/Time"),
    ("movement_type", "Type"),
    ("product_name", "Product"),
    ("product_code", "Product Code"),
    ("lot_number", "Lot Number"),
    ("quantity", "Quantity"),
    ("user", "User"),
    ("location_name", "Location"),
]


def _parse_date(value, end_of_day=False):
    if not value:
        return None
    parsed = datetime.datetime.strptime(value, "%Y-%m-%d")
    if end_of_day:
        parsed = parsed.replace(hour=23, minute=59, second=59, microsecond=999999)
    return make_aware(parsed)


def _apply_date_filter(qs, start_date, end_date, field="timestamp"):
    if start_date:
        qs = qs.filter(**{f"{field}__gte": _parse_date(start_date)})
    if end_date:
        qs = qs.filter(**{f"{field}__lte": _parse_date(end_date, end_of_day=True)})
    return qs


def _format_value(value):
    if value is None:
        return ""
    if isinstance(value, datetime.datetime):
        return value.strftime("%Y-%m-%d %H:%M")
    if isinstance(value, datetime.date):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, Decimal):
        # Trim trailing zeros so 13.00 -> "13", 4.50 -> "4.5".
        return format(value.normalize(), "f")
    return str(value)


def _involved_item_ids(start_date, end_date):
    """Product items (lots) that had ANY action in the range: an arrival
    (registration), a withdrawal, or a QC check."""
    reg = _apply_date_filter(StockRegistration.objects.all(), start_date, end_date)
    wd = _apply_date_filter(Withdrawal.objects.all(), start_date, end_date)
    ids = set(reg.values_list("product_item_id", flat=True))
    ids |= set(wd.values_list("product_item_id", flat=True))
    if QualityCheck is not None:
        qc = _apply_date_filter(
            QualityCheck.objects.all(), start_date, end_date, field="created_at"
        )
        ids |= set(qc.values_list("product_item_id", flat=True))
    ids.discard(None)
    return ids


def _build_lot_summary_rows(item_ids):
    """One row per lot. Received/Withdrawn are lifetime totals for the lot so
    they reconcile with the lot's current stock (a single withdrawal shows as 1)."""
    if not item_ids:
        return []

    today = timezone.localdate()
    soon = today + datetime.timedelta(days=30)

    received = defaultdict(int)
    first_reg = {}
    for r in (
        StockRegistration.objects.filter(product_item_id__in=item_ids)
        .order_by("timestamp", "id")
        .values("product_item_id", "quantity", "timestamp", "user__username")
    ):
        pid = r["product_item_id"]
        received[pid] += int(r["quantity"] or 0)
        if pid not in first_reg:
            first_reg[pid] = (r["timestamp"], r["user__username"])

    withdrawn = defaultdict(Decimal)
    last_wd = {}
    for w in (
        Withdrawal.objects.filter(product_item_id__in=item_ids)
        .order_by("timestamp", "id")
        .values("product_item_id", "quantity", "timestamp", "user__username")
    ):
        pid = w["product_item_id"]
        withdrawn[pid] += Decimal(w["quantity"] or 0)
        last_wd[pid] = (w["timestamp"], w["user__username"])

    qc_latest = {}
    if QualityCheck is not None:
        for q in (
            QualityCheck.objects.filter(product_item_id__in=item_ids)
            .order_by("created_at", "id")
            .values("product_item_id", "result", "performed_by__username", "created_at")
        ):
            qc_latest[q["product_item_id"]] = (
                q["result"],
                q["performed_by__username"],
                q["created_at"],
            )

    items = (
        ProductItem.objects.filter(id__in=item_ids)
        .select_related("product", "product__location")
        .order_by("product__name", "lot_number", "expiry_date")
    )

    rows = []
    for it in items:
        pid = it.id
        result, qc_user, qc_ts = qc_latest.get(pid, (None, None, None))
        if result == "pass":
            qc_status = "Passed"
        elif result == "fail":
            qc_status = "Failed"
        else:
            qc_status = "Waiting"

        if it.expiry_date and it.expiry_date <= today:
            expiry_status = "Expired"
        elif it.expiry_date and it.expiry_date <= soon:
            expiry_status = "Expiring soon"
        else:
            expiry_status = "OK"

        fr_ts, fr_user = first_reg.get(pid, (None, None))
        lw_ts, lw_user = last_wd.get(pid, (None, None))
        product = it.product

        rows.append(
            {
                "product_name": product.name if product else "",
                "product_code": (product.product_code if product else "") or "",
                "lot_number": it.lot_number,
                "expiry_date": it.expiry_date,
                "expiry_status": expiry_status,
                "location_name": product.location.name if product and product.location else "",
                "received": received.get(pid, 0),
                "withdrawn": withdrawn.get(pid, Decimal(0)),
                "current_stock": it.current_stock,
                "qc_status": qc_status,
                "qc_by": qc_user or "",
                "qc_date": qc_ts,
                "first_registered": fr_ts,
                "registered_by": fr_user or "",
                "last_withdrawal": lw_ts,
                "withdrawn_by": lw_user or "",
            }
        )
    return rows


def _build_movements_rows(start_date, end_date):
    """Every action in range as its own row (audit trail): registrations,
    withdrawals, and QC checks."""
    movements = []

    reg = _apply_date_filter(
        StockRegistration.objects.select_related("user", "location"), start_date, end_date
    )
    for r in reg.values(
        "timestamp", "product_name", "product_code", "lot_number",
        "quantity", "user__username", "location__name",
    ):
        movements.append(
            {
                "movement_date": r["timestamp"],
                "movement_type": "Registered (in)",
                "product_name": r["product_name"],
                "product_code": r["product_code"],
                "lot_number": r["lot_number"],
                "quantity": r["quantity"],
                "user": r["user__username"] or "",
                "location_name": r["location__name"] or "",
            }
        )

    wd = _apply_date_filter(
        Withdrawal.objects.select_related("user", "location"), start_date, end_date
    )
    for w in wd.values(
        "timestamp", "product_name", "product_code", "lot_number",
        "quantity", "user__username", "location__name",
    ):
        movements.append(
            {
                "movement_date": w["timestamp"],
                "movement_type": "Withdrawn (out)",
                "product_name": w["product_name"],
                "product_code": w["product_code"],
                "lot_number": w["lot_number"],
                "quantity": w["quantity"],
                "user": w["user__username"] or "",
                "location_name": w["location__name"] or "",
            }
        )

    if QualityCheck is not None:
        qc = _apply_date_filter(
            QualityCheck.objects.select_related(
                "product_item__product", "product_item__product__location",
                "performed_by", "location",
            ),
            start_date, end_date, field="created_at",
        )
        for q in qc.values(
            "created_at", "result",
            "product_item__product__name", "product_item__product__product_code",
            "product_item__lot_number", "performed_by__username",
            "location__name", "product_item__product__location__name",
        ):
            result = q["result"]
            label = "Pass" if result == "pass" else ("Fail" if result == "fail" else "Pending")
            movements.append(
                {
                    "movement_date": q["created_at"],
                    "movement_type": f"QC check ({label})",
                    "product_name": q["product_item__product__name"] or "",
                    "product_code": q["product_item__product__product_code"] or "",
                    "lot_number": q["product_item__lot_number"] or "",
                    "quantity": None,
                    "user": q["performed_by__username"] or "",
                    "location_name": q["location__name"] or q["product_item__product__location__name"] or "",
                }
            )

    movements.sort(key=lambda m: m["movement_date"], reverse=True)
    return movements


def _write_sheet(ws, columns, rows):
    headers = [label for _, label in columns]
    ws.append(headers)
    for row in rows:
        ws.append([_format_value(row[key]) for key, _ in columns])
    ws.freeze_panes = "A2"
    if rows:
        ws.auto_filter.ref = ws.dimensions
    for idx, (_, label) in enumerate(columns, start=1):
        ws.column_dimensions[ws.cell(row=1, column=idx).column_letter].width = min(
            max(len(label) + 2, 12), 42
        )


def _download_response(summary_rows, movements_rows, download_type, user_part, date_part):
    if download_type == "excel":
        wb = Workbook()
        summary_ws = wb.active
        summary_ws.title = "Lot Summary"
        _write_sheet(summary_ws, LOT_SUMMARY_COLUMNS, summary_rows)

        movements_ws = wb.create_sheet(title="Activity Log")
        _write_sheet(movements_ws, MOVEMENTS_COLUMNS, movements_rows)

        out = io.BytesIO()
        wb.save(out)
        out.seek(0)
        response = HttpResponse(
            out,
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        response["Content-Disposition"] = f"attachment; filename=IMS_{user_part}_{date_part}.xlsx"
        return response

    # CSV: single sheet only -> emit the lot summary (the primary view).
    csv_buffer = io.StringIO()
    writer = csv.writer(csv_buffer)
    writer.writerow([label for _, label in LOT_SUMMARY_COLUMNS])
    for row in summary_rows:
        writer.writerow([_format_value(row[key]) for key, _ in LOT_SUMMARY_COLUMNS])
    response = HttpResponse(csv_buffer.getvalue(), content_type="text/csv")
    response["Content-Disposition"] = f"attachment; filename=IMS_{user_part}_{date_part}.csv"
    return response


def download_report(request):
    start_date = request.GET.get("start_date")
    end_date = request.GET.get("end_date")
    download_type = request.GET.get("download")
    page_number = request.GET.get("page", 1)

    item_ids = _involved_item_ids(start_date, end_date)
    summary_rows = _build_lot_summary_rows(item_ids)

    if download_type in {"excel", "csv"}:
        movements_rows = _build_movements_rows(start_date, end_date)
        user_part = request.user.username if request.user.is_authenticated else "anonymous"
        date_part = now().strftime("%Y%m%d")
        return _download_response(summary_rows, movements_rows, download_type, user_part, date_part)

    paginator = Paginator(summary_rows, 100)
    page_obj = paginator.get_page(page_number)

    return render(
        request,
        "analytics/download_report.html",
        {
            "columns": LOT_SUMMARY_COLUMNS,
            "rows": page_obj.object_list,
            "page_obj": page_obj,
            "start_date": start_date,
            "end_date": end_date,
        },
    )
